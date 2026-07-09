"""Build a structured DOCX transcript document from timestamped segments."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .transcriber import TranscriptSegment

_PARAGRAPH_GAP_SECONDS = 2.0
_SECTION_GAP_SECONDS = 8.0


def _format_timestamp(seconds: float) -> str:
    total = int(seconds)
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    if hours:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"
    return f"{minutes:02d}:{secs:02d}"


def _group_segments(segments: list[TranscriptSegment]) -> list[list[list[TranscriptSegment]]]:
    """Group segments into sections of paragraphs based on pause length."""
    sections: list[list[list[TranscriptSegment]]] = []
    current_section: list[list[TranscriptSegment]] = []
    current_paragraph: list[TranscriptSegment] = []
    prev_end = None

    for seg in segments:
        if prev_end is not None:
            gap = seg.start - prev_end
            if gap >= _SECTION_GAP_SECONDS:
                if current_paragraph:
                    current_section.append(current_paragraph)
                    current_paragraph = []
                if current_section:
                    sections.append(current_section)
                    current_section = []
            elif gap >= _PARAGRAPH_GAP_SECONDS and current_paragraph:
                current_section.append(current_paragraph)
                current_paragraph = []
        current_paragraph.append(seg)
        prev_end = seg.end

    if current_paragraph:
        current_section.append(current_paragraph)
    if current_section:
        sections.append(current_section)

    return sections


def build_document(
    segments: list[TranscriptSegment],
    source_url: str,
    title: str,
    out_path: Path,
) -> Path:
    doc = Document()

    heading = doc.add_heading(title, level=0)
    heading.alignment = 0

    meta = doc.add_paragraph()
    meta.add_run("Source: ").bold = True
    meta.add_run(source_url)
    meta.add_run("\n")
    meta.add_run("Generated: ").bold = True
    meta.add_run(datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    if segments:
        meta.add_run("\n")
        meta.add_run("Duration: ").bold = True
        meta.add_run(_format_timestamp(segments[-1].end))

    doc.add_paragraph()  # spacer

    sections = _group_segments(segments)
    for i, section in enumerate(sections, start=1):
        section_start = section[0][0].start
        doc.add_heading(f"Section {i} — {_format_timestamp(section_start)}", level=1)

        for paragraph_segments in section:
            para_start = paragraph_segments[0].start
            text = " ".join(seg.text for seg in paragraph_segments)

            p = doc.add_paragraph()
            timestamp_run = p.add_run(f"[{_format_timestamp(para_start)}] ")
            timestamp_run.bold = True
            timestamp_run.font.size = Pt(10)
            text_run = p.add_run(text)
            text_run.font.size = Pt(11)

    if not sections:
        doc.add_paragraph("No speech was detected in this video.")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
